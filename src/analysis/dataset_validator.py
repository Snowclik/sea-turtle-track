"""
Etapa 2.5 — Dataset Validator
src/analysis/dataset_validator.py

Valida la integridad completa del dataset original sin modificar ningún archivo.
Genera reportes en HTML, PDF, JSON, Excel y previews visuales por imagen.
"""

import os
import glob
import json
import hashlib
import csv

import cv2
import numpy as np
from PIL import Image
from PIL.ExifTags import TAGS
from jinja2 import Template
import openpyxl
from rich.console import Console
from rich.table import Table
from rich.panel import Panel
from rich.progress import track
from rich import box
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table as RLTable, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors as rl_colors

# ─── Configuración ────────────────────────────────────────────────────────────

console = Console(force_terminal=True, highlight=False)

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
DATASET_IMG_DIR = os.path.join(BASE_DIR, "dataset", "original")
DATASET_LBL_DIR = os.path.join(BASE_DIR, "dataset", "labels")
VALIDATION_DIR  = os.path.join(BASE_DIR, "reports", "validation")
PREVIEWS_DIR    = os.path.join(VALIDATION_DIR, "previews")

VALID_CLASSES   = {0}          # 0 → nesting_track
VALID_IMG_EXT   = ('.jpg', '.jpeg', '.png')

os.makedirs(PREVIEWS_DIR, exist_ok=True)


# ─── Clase Principal ──────────────────────────────────────────────────────────

class DatasetValidator:

    def __init__(self):
        self.images = sorted([
            p for p in glob.glob(os.path.join(DATASET_IMG_DIR, "*.*"))
            if p.lower().endswith(VALID_IMG_EXT)
        ])
        self.labels = sorted(glob.glob(os.path.join(DATASET_LBL_DIR, "*.txt")))

        # Mapas nombre → ruta
        self.img_map = {os.path.splitext(os.path.basename(p))[0]: p for p in self.images}
        self.lbl_map = {os.path.splitext(os.path.basename(p))[0]: p for p in self.labels}

        # Resultados de validación
        self.errors   = []          # lista de dict {type, file, detail}
        self.warnings = []
        self.manifest = []          # un entry por imagen

        # Contadores globales
        self.total_objects    = 0
        self.valid_images     = 0
        self.valid_labels     = 0
        self.valid_bboxes     = 0
        self.total_bboxes     = 0
        self.corrupted_images = 0
        self.missing_labels   = 0
        self.invalid_boxes    = 0
        self.invalid_classes  = 0
        self.empty_labels     = 0
        self.dqi              = 0.0

    # ── Validaciones ──────────────────────────────────────────────────────────

    def validate_images(self):
        console.print("\n[bold cyan]>> Validando imagenes...[/bold cyan]")

        hashes = {}
        for img_path in track(self.images, description="  Imágenes"):
            name = os.path.splitext(os.path.basename(img_path))[0]
            issues = []

            # Archivo vacío
            if os.path.getsize(img_path) == 0:
                issues.append("archivo vacío")
                self.errors.append({"type": "empty_image", "file": img_path, "detail": "Archivo vacío"})
                self.corrupted_images += 1
                self._add_manifest(name, img_path, issues=issues, status="invalid")
                continue

            # Corrupción
            img = cv2.imread(img_path)
            if img is None:
                issues.append("imagen corrupta")
                self.errors.append({"type": "corrupted", "file": img_path, "detail": "No se puede leer"})
                self.corrupted_images += 1
                self._add_manifest(name, img_path, issues=issues, status="invalid")
                continue

            # Duplicados (hash MD5)
            with open(img_path, 'rb') as f:
                fhash = hashlib.md5(f.read()).hexdigest()
            if fhash in hashes:
                issues.append(f"duplicada de {os.path.basename(hashes[fhash])}")
                self.warnings.append({"type": "duplicate", "file": img_path, "detail": f"Duplicada de {hashes[fhash]}"})
            else:
                hashes[fhash] = img_path

            # Sin etiqueta
            if name not in self.lbl_map:
                issues.append("sin etiqueta")
                self.missing_labels += 1

            self.valid_images += 1
            self._add_manifest(name, img_path, issues=issues, status="valid" if not issues else "warning")

    def validate_labels(self):
        console.print("[bold cyan]>> Validando etiquetas...[/bold cyan]")

        for lbl_path in track(self.labels, description="  Etiquetas"):
            name = os.path.splitext(os.path.basename(lbl_path))[0]

            # Etiqueta sin imagen
            if name not in self.img_map:
                self.errors.append({"type": "orphan_label", "file": lbl_path, "detail": "Sin imagen correspondiente"})
                continue

            # Archivo vacío
            if os.path.getsize(lbl_path) == 0:
                self.empty_labels += 1
                self.warnings.append({"type": "empty_label", "file": lbl_path, "detail": "Archivo vacío"})
                self._update_manifest(name, label_path=lbl_path, objects=0, valid_label=True)
                continue

            # Parsear líneas
            try:
                with open(lbl_path, 'r') as f:
                    lines = [l.strip() for l in f.readlines() if l.strip()]
            except Exception as e:
                self.errors.append({"type": "read_error", "file": lbl_path, "detail": str(e)})
                continue

            bbox_list = []
            file_ok = True
            for i, line in enumerate(lines, 1):
                parts = line.split()

                # Mínimo: clase + 4 coords (bbox) o clase + 6 coords (polígono mínimo)
                if len(parts) < 5:
                    self.errors.append({"type": "format_error", "file": lbl_path,
                                        "detail": f"Linea {i}: {len(parts)} campos (minimo 5)"})
                    file_ok = False
                    self.invalid_boxes += 1
                    self.total_bboxes += 1
                    continue

                # Verificar que todos sean numéricos
                try:
                    cls = int(parts[0])
                    coords = [float(v) for v in parts[1:]]
                except ValueError:
                    self.errors.append({"type": "format_error", "file": lbl_path,
                                        "detail": f"Linea {i}: valores no numericos"})
                    file_ok = False
                    self.invalid_boxes += 1
                    self.total_bboxes += 1
                    continue

                self.total_bboxes += 1
                box_ok = True

                # Detectar tipo: bbox (4 coords) o segmentación (N pares >= 3)
                n_coords = len(coords)
                is_bbox = (n_coords == 4)
                is_seg  = (n_coords >= 6 and n_coords % 2 == 0)

                if not is_bbox and not is_seg:
                    self.errors.append({"type": "format_error", "file": lbl_path,
                                        "detail": f"Linea {i}: {n_coords} coordenadas (se esperan 4 o N pares)"})
                    file_ok = False
                    self.invalid_boxes += 1
                    self.total_bboxes -= 1
                    continue

                # Clase inválida
                if cls not in VALID_CLASSES:
                    self.errors.append({"type": "invalid_class", "file": lbl_path,
                                        "detail": f"Linea {i}: clase {cls} no existe"})
                    self.invalid_classes += 1
                    box_ok = False

                # Coordenadas fuera de rango [0, 1]
                for j, val in enumerate(coords):
                    if val < 0:
                        self.errors.append({"type": "negative_coord", "file": lbl_path,
                                            "detail": f"Linea {i}: coord[{j}]={val:.4f} < 0"})
                        box_ok = False
                    elif val > 1:
                        self.errors.append({"type": "out_of_range", "file": lbl_path,
                                            "detail": f"Linea {i}: coord[{j}]={val:.4f} > 1"})
                        box_ok = False

                # BBox: validar w/h > 0
                if is_bbox:
                    bw, bh = coords[2], coords[3]
                    if bw == 0 or bh == 0:
                        self.errors.append({"type": "zero_dimension", "file": lbl_path,
                                            "detail": f"Linea {i}: bw={bw} bh={bh}"})
                        box_ok = False
                    if box_ok:
                        bbox_list.append((cls, coords[0], coords[1], bw, bh, "bbox"))
                else:
                    # Segmentación: guardar como (cls, xs, ys, "seg")
                    xs = coords[0::2]
                    ys = coords[1::2]
                    if box_ok:
                        bbox_list.append((cls, xs, ys, "seg"))

                if box_ok:
                    self.valid_bboxes += 1
                else:
                    self.invalid_boxes += 1

            n_obj = len(lines)
            self.total_objects += n_obj
            self.valid_labels += 1 if file_ok else 0
            self._update_manifest(name, label_path=lbl_path, objects=n_obj,
                                  valid_label=file_ok, bboxes=bbox_list)

    # ── DQI ──────────────────────────────────────────────────────────────────

    def compute_dqi(self):
        console.print("[bold cyan]>> Calculando Dataset Quality Index...[/bold cyan]")
        n = len(self.images)
        if n == 0:
            self.dqi = 0.0
            return

        # Componentes
        w_valid_img  = 0.25
        w_label_cov  = 0.25
        w_bbox_valid = 0.20
        w_cls_valid  = 0.15
        w_general    = 0.15

        c_valid_img  = (self.valid_images   / n) if n else 0
        c_label_cov  = ((n - self.missing_labels) / n) if n else 0
        c_bbox       = (self.valid_bboxes   / self.total_bboxes) if self.total_bboxes else 1.0
        c_cls        = 1.0 - (self.invalid_classes / self.total_bboxes) if self.total_bboxes else 1.0
        c_general    = 1.0 - (len(self.errors) / max(1, n * 3))
        c_general    = max(0.0, c_general)

        self.dqi = round(
            (c_valid_img * w_valid_img +
             c_label_cov * w_label_cov +
             c_bbox      * w_bbox_valid +
             c_cls       * w_cls_valid +
             c_general   * w_general) * 100, 2
        )

    # ── Previews ──────────────────────────────────────────────────────────────

    def generate_previews(self):
        console.print("[bold cyan]>> Generando previews...[/bold cyan]")
        COLORS = [(0, 230, 64), (255, 80, 80), (80, 160, 255)]

        for img_path in track(self.images, description="  Previews"):
            name = os.path.splitext(os.path.basename(img_path))[0]
            img  = cv2.imread(img_path)
            if img is None:
                continue

            h, w, _ = img.shape
            n_objects = 0

            if name in self.lbl_map:
                lbl_path = self.lbl_map[name]
                try:
                    with open(lbl_path, 'r') as f:
                        for line in f:
                            parts = line.strip().split()
                            if len(parts) < 5:
                                continue
                            cls_id = int(parts[0])
                            coords = [float(v) for v in parts[1:]]
                            color  = COLORS[cls_id % len(COLORS)]
                            n_coords = len(coords)

                            if n_coords == 4:
                                # BBox estándar
                                cx, cy, bw, bh = coords
                                x1 = int((cx - bw / 2) * w)
                                y1 = int((cy - bh / 2) * h)
                                x2 = int((cx + bw / 2) * w)
                                y2 = int((cy + bh / 2) * h)
                                cv2.rectangle(img, (x1, y1), (x2, y2), color, 3)
                                cv2.putText(img, f"nesting_track", (x1, max(y1 - 10, 20)),
                                            cv2.FONT_HERSHEY_SIMPLEX, 0.7, color, 2)
                            elif n_coords >= 6 and n_coords % 2 == 0:
                                # Segmentación polígono
                                xs = [int(v * w) for v in coords[0::2]]
                                ys = [int(v * h) for v in coords[1::2]]
                                pts = np.array(list(zip(xs, ys)), dtype=np.int32)
                                cv2.polylines(img, [pts], isClosed=True, color=color, thickness=3)
                                # Overlay translúcido del polígono
                                overlay = img.copy()
                                cv2.fillPoly(overlay, [pts], color)
                                img = cv2.addWeighted(overlay, 0.25, img, 0.75, 0)
                                # Etiqueta en el centroide
                                cx_px = int(np.mean(xs))
                                cy_px = int(np.mean(ys))
                                cv2.putText(img, "nesting_track",
                                            (max(cx_px - 50, 0), max(cy_px - 10, 20)),
                                            cv2.FONT_HERSHEY_SIMPLEX, 0.6, color, 2)
                            n_objects += 1
                except Exception:
                    pass

            # Info overlay
            info = f"{os.path.basename(img_path)} | {w}x{h} | Objetos: {n_objects}"
            cv2.rectangle(img, (0, 0), (w, 40), (0, 0, 0), -1)
            cv2.putText(img, info, (10, 28), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2)

            if w > 1280:
                scale = 1280 / w
                img = cv2.resize(img, (1280, int(h * scale)))

            out_path = os.path.join(PREVIEWS_DIR, f"preview_{os.path.basename(img_path)}")
            cv2.imwrite(out_path, img)

    # ── Manifest ──────────────────────────────────────────────────────────────

    def generate_manifest(self):
        console.print("[bold cyan]>> Generando dataset manifest...[/bold cyan]")
        manifest_path = os.path.join(VALIDATION_DIR, "dataset_manifest.json")
        with open(manifest_path, 'w', encoding='utf-8') as f:
            json.dump(self.manifest, f, indent=4, ensure_ascii=False)

    # ── Helpers de manifest ───────────────────────────────────────────────────

    def _add_manifest(self, name, img_path, issues=None, status="valid"):
        self.manifest.append({
            "image":   os.path.basename(img_path),
            "label":   os.path.basename(self.lbl_map.get(name, "")) or "N/A",
            "objects": 0,
            "class":   "nesting_track",
            "status":  status,
            "issues":  issues or []
        })

    def _update_manifest(self, name, label_path=None, objects=0, valid_label=True, bboxes=None):
        for entry in self.manifest:
            img_base = os.path.splitext(entry["image"])[0]
            if img_base == name:
                if label_path:
                    entry["label"]   = os.path.basename(label_path)
                entry["objects"] = objects
                entry["status"]  = "valid" if valid_label and not entry["issues"] else "warning"
                return
        # Si no existe la entrada (imagen no estaba), agregar
        img_path = self.img_map.get(name, name)
        self.manifest.append({
            "image":   os.path.basename(img_path),
            "label":   os.path.basename(label_path) if label_path else "N/A",
            "objects": objects,
            "class":   "nesting_track",
            "status":  "valid" if valid_label else "invalid",
            "issues":  []
        })

    # ── Reportes ──────────────────────────────────────────────────────────────

    def generate_reports(self):
        console.print("[bold cyan]>> Generando reportes...[/bold cyan]")

        n  = len(self.images)
        n_labeled = n - self.missing_labels
        pct_labeled = round((n_labeled / n) * 100, 2) if n else 0
        ready = "Sí" if self.dqi >= 95 else "No"

        summary = {
            "total_imagenes":       n,
            "total_etiquetas":      len(self.labels),
            "total_objetos":        self.total_objects,
            "total_bboxes":         self.total_bboxes,
            "imagenes_corruptas":   self.corrupted_images,
            "etiquetas_faltantes":  self.missing_labels,
            "bboxes_invalidos":     self.invalid_boxes,
            "errores_totales":      len(self.errors),
            "advertencias":         len(self.warnings),
            "pct_imagenes_con_lbl": pct_labeled,
            "dqi":                  self.dqi,
            "listo_para_entreno":   ready,
            "errores":              self.errors[:50],   # primeros 50
            "advertencias_detalle": self.warnings[:20]
        }

        # JSON
        json_path = os.path.join(VALIDATION_DIR, "validation_summary.json")
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(summary, f, indent=4, ensure_ascii=False)

        # Excel
        wb = openpyxl.Workbook()
        ws_sum = wb.active
        ws_sum.title = "Resumen"
        ws_sum.append(["Métrica", "Valor"])
        for k, v in summary.items():
            if not isinstance(v, list):
                ws_sum.append([k, str(v)])

        ws_mfst = wb.create_sheet("Manifest")
        ws_mfst.append(["image", "label", "objects", "class", "status", "issues"])
        for entry in self.manifest:
            ws_mfst.append([
                entry["image"], entry["label"], entry["objects"],
                entry["class"], entry["status"], "; ".join(entry["issues"])
            ])

        ws_err = wb.create_sheet("Errores")
        ws_err.append(["type", "file", "detail"])
        for e in self.errors:
            ws_err.append([e["type"], e["file"], e["detail"]])

        wb.save(os.path.join(VALIDATION_DIR, "validation_summary.xlsx"))

        # HTML
        self._generate_html(summary)

        # PDF
        self._generate_pdf(summary)

    def _generate_html(self, summary):
        preview_files = sorted([
            f for f in os.listdir(PREVIEWS_DIR)
            if f.lower().endswith(('.jpg', '.jpeg', '.png'))
        ])

        template_str = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>Validación Dataset ASFP</title>
    <style>
        body { font-family: 'Segoe UI', Arial, sans-serif; background:#0f1117; color:#e0e0e0; margin:0; padding:30px; }
        h1 { color:#00e6a0; letter-spacing:2px; }
        h2 { color:#60b4ff; border-bottom:1px solid #333; padding-bottom:6px; }
        .dqi { font-size:3rem; font-weight:bold; color:{{ dqi_color }}; }
        table { border-collapse:collapse; width:100%; margin-bottom:28px; }
        th { background:#1e2130; color:#60b4ff; padding:10px; text-align:left; }
        td { padding:8px 10px; border-bottom:1px solid #2a2d3a; }
        tr:hover td { background:#1a1d28; }
        .badge-valid   { background:#1a4a2e; color:#00e6a0; padding:2px 8px; border-radius:4px; }
        .badge-warning { background:#4a3a0a; color:#ffc107; padding:2px 8px; border-radius:4px; }
        .badge-invalid { background:#4a1a1a; color:#ff4444; padding:2px 8px; border-radius:4px; }
        .grid { display:flex; flex-wrap:wrap; gap:14px; }
        .preview-card { background:#1a1d28; border:1px solid #2a2d3a; border-radius:8px; padding:8px; width:280px; }
        .preview-card img { width:100%; border-radius:4px; }
        .preview-card p { font-size:0.75rem; color:#888; margin:4px 0 0; text-align:center; }
        .stat-grid { display:grid; grid-template-columns:repeat(3,1fr); gap:14px; margin-bottom:28px; }
        .stat-card { background:#1a1d28; border:1px solid #2a2d3a; border-radius:8px; padding:16px; }
        .stat-val { font-size:1.6rem; font-weight:bold; color:#00e6a0; }
        .stat-label { font-size:0.8rem; color:#888; margin-top:4px; }
        .error-row td { color:#ff6b6b; }
        .warn-row td  { color:#ffc107; }
    </style>
</head>
<body>
<h1>📊 Dataset Validation Report — ASFP</h1>
<div style="display:flex;align-items:center;gap:30px;margin-bottom:30px;">
    <div>
        <div style="font-size:0.9rem;color:#888;margin-bottom:4px;">Dataset Quality Index</div>
        <div class="dqi">{{ dqi }}%</div>
        <div style="color:{% if ready == 'Sí' %}#00e6a0{% else %}#ff4444{% endif %};font-size:1rem;margin-top:4px;">
            Listo para Entrenamiento: <strong>{{ ready }}</strong>
        </div>
    </div>
</div>

<div class="stat-grid">
    <div class="stat-card"><div class="stat-val">{{ total_imagenes }}</div><div class="stat-label">Imágenes Totales</div></div>
    <div class="stat-card"><div class="stat-val">{{ total_etiquetas }}</div><div class="stat-label">Etiquetas Totales</div></div>
    <div class="stat-card"><div class="stat-val">{{ total_objetos }}</div><div class="stat-label">Objetos Totales</div></div>
    <div class="stat-card"><div class="stat-val">{{ imagenes_corruptas }}</div><div class="stat-label">Imágenes Corruptas</div></div>
    <div class="stat-card"><div class="stat-val">{{ etiquetas_faltantes }}</div><div class="stat-label">Etiquetas Faltantes</div></div>
    <div class="stat-card"><div class="stat-val">{{ bboxes_invalidos }}</div><div class="stat-label">BBoxes Inválidos</div></div>
</div>

<h2>Manifest del Dataset</h2>
<table>
<thead><tr><th>Imagen</th><th>Etiqueta</th><th>Objetos</th><th>Estado</th><th>Problemas</th></tr></thead>
<tbody>
{% for entry in manifest %}
<tr>
    <td>{{ entry.image }}</td>
    <td>{{ entry.label }}</td>
    <td>{{ entry.objects }}</td>
    <td><span class="badge-{{ entry.status }}">{{ entry.status }}</span></td>
    <td style="color:#ffc107;font-size:0.8rem;">{{ entry.issues | join(', ') }}</td>
</tr>
{% endfor %}
</tbody>
</table>

{% if errors %}
<h2>⚠ Errores Detectados</h2>
<table>
<thead><tr><th>Tipo</th><th>Archivo</th><th>Detalle</th></tr></thead>
<tbody>
{% for e in errors %}
<tr class="error-row"><td>{{ e.type }}</td><td>{{ e.file }}</td><td>{{ e.detail }}</td></tr>
{% endfor %}
</tbody>
</table>
{% endif %}

<h2>🖼 Previews</h2>
<div class="grid">
{% for fname in preview_files %}
<div class="preview-card">
    <img src="./previews/{{ fname }}" alt="{{ fname }}">
    <p>{{ fname }}</p>
</div>
{% endfor %}
</div>
</body>
</html>"""

        from jinja2 import Template
        dqi_color = "#00e6a0" if summary["dqi"] >= 95 else ("#ffc107" if summary["dqi"] >= 70 else "#ff4444")
        html = Template(template_str).render(
            dqi=summary["dqi"],
            dqi_color=dqi_color,
            ready=summary["listo_para_entreno"],
            total_imagenes=summary["total_imagenes"],
            total_etiquetas=summary["total_etiquetas"],
            total_objetos=summary["total_objetos"],
            imagenes_corruptas=summary["imagenes_corruptas"],
            etiquetas_faltantes=summary["etiquetas_faltantes"],
            bboxes_invalidos=summary["bboxes_invalidos"],
            manifest=self.manifest,
            errors=self.errors,
            preview_files=preview_files
        )
        with open(os.path.join(VALIDATION_DIR, "validation_report.html"), 'w', encoding='utf-8') as f:
            f.write(html)

    def _generate_pdf(self, summary):
        pdf_path = os.path.join(VALIDATION_DIR, "validation_report.pdf")
        doc = SimpleDocTemplate(pdf_path, pagesize=letter)
        styles = getSampleStyleSheet()
        elems = []

        elems.append(Paragraph("Dataset Validation Report — ASFP", styles['Title']))
        elems.append(Spacer(1, 12))
        elems.append(Paragraph(f"Dataset Quality Index (DQI): {summary['dqi']}%", styles['Heading2']))
        elems.append(Paragraph(f"Listo para entrenamiento: {summary['listo_para_entreno']}", styles['Normal']))
        elems.append(Spacer(1, 12))

        data = [
            ["Métrica", "Valor"],
            ["Total Imágenes",       str(summary["total_imagenes"])],
            ["Total Etiquetas",      str(summary["total_etiquetas"])],
            ["Total Objetos",        str(summary["total_objetos"])],
            ["Imágenes Corruptas",   str(summary["imagenes_corruptas"])],
            ["Etiquetas Faltantes",  str(summary["etiquetas_faltantes"])],
            ["BBoxes Inválidos",     str(summary["bboxes_invalidos"])],
            ["Errores Totales",      str(summary["errores_totales"])],
            ["% Etiquetado",         f"{summary['pct_imagenes_con_lbl']}%"],
        ]
        t = RLTable(data, colWidths=[280, 120])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), rl_colors.HexColor('#2c3e50')),
            ('TEXTCOLOR',  (0, 0), (-1, 0), rl_colors.white),
            ('ALIGN',      (0, 0), (-1, -1), 'CENTER'),
            ('GRID',       (0, 0), (-1, -1), 0.5, rl_colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [rl_colors.white, rl_colors.HexColor('#f9f9f9')]),
        ]))
        elems.append(t)
        elems.append(Spacer(1, 16))

        if self.errors:
            elems.append(Paragraph("Errores Detectados", styles['Heading2']))
            err_data = [["Tipo", "Archivo", "Detalle"]]
            for e in self.errors[:30]:
                err_data.append([e["type"], os.path.basename(e["file"]), e["detail"][:60]])
            et = RLTable(err_data, colWidths=[120, 160, 140])
            et.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), rl_colors.HexColor('#c0392b')),
                ('TEXTCOLOR',  (0, 0), (-1, 0), rl_colors.white),
                ('GRID',       (0, 0), (-1, -1), 0.5, rl_colors.grey),
                ('FONTSIZE',   (0, 0), (-1, -1), 7),
            ]))
            elems.append(et)

        doc.build(elems)

    # ── Resumen en consola ─────────────────────────────────────────────────────

    def print_summary(self):
        n = len(self.images)
        dqi_color = "green" if self.dqi >= 95 else ("yellow" if self.dqi >= 70 else "red")

        table = Table(title="DATASET VALIDATION COMPLETED", box=box.DOUBLE_EDGE,
                      border_style="cyan", show_header=False, padding=(0, 2))
        table.add_column("Métrica",  style="bold white", width=30)
        table.add_column("Valor",    style="bold", width=20)

        table.add_row("Images",           str(n))
        table.add_row("Labels",           str(len(self.labels)))
        table.add_row("Objects",          str(self.total_objects))
        table.add_row("Corrupted Images", str(self.corrupted_images))
        table.add_row("Missing Labels",   str(self.missing_labels))
        table.add_row("Invalid Boxes",    str(self.invalid_boxes))
        table.add_row("Total Errors",     str(len(self.errors)))
        table.add_row("Dataset Quality",  f"[{dqi_color}]{self.dqi} %[/{dqi_color}]")
        table.add_row("HTML Report",      "[green]OK[/green]")
        table.add_row("PDF Report",       "[green]OK[/green]")
        table.add_row("Preview Images",   "[green]OK[/green]")

        console.print()
        console.print(table)

        status_msg  = "[bold green]✔ Dataset listo para Etapa 3[/bold green]" \
                      if self.dqi >= 95 else \
                      "[bold yellow]⚠ Corrige los errores antes de continuar[/bold yellow]"
        console.print(Panel(status_msg, border_style=dqi_color))


# ── Entrada directa ───────────────────────────────────────────────────────────

def run_validation():
    validator = DatasetValidator()
    validator.validate_images()
    validator.validate_labels()
    validator.compute_dqi()
    validator.generate_previews()
    validator.generate_manifest()
    validator.generate_reports()
    validator.print_summary()
    return validator
